#!/usr/bin/env python3
"""
Template Validation Utility for StyleStack

This tool validates OOXML template integrity and identifies corruption issues.
IMPORTANT: Run with activated venv - this is Brew Python!

Usage:
    source venv/bin/activate
    python tools/template_validator.py --template path/to/template.potx
"""


from typing import Any, Dict, List
import argparse
import sys
import zipfile
from pathlib import Path
import xml.etree.ElementTree as ET
from xml.parsers.expat import ExpatError

# PyOffice imports - requires venv activation
try:
    from docx import Document
    from openpyxl import load_workbook
    from pptx import Presentation
    import lxml.etree as lxml_ET
    from lxml.etree import XMLSyntaxError
except ImportError as e:
    print(f"ERROR: PyOffice dependencies not found. Did you activate venv?")
    print(f"Run: source venv/bin/activate")
    print(f"Missing: {e}")
    sys.exit(1)


class TemplateValidator:
    """Validates OOXML templates for corruption and structure issues"""
    
    def __init__(self, template_path: str):
        self.template_path = Path(template_path)
        self.issues: List[Dict[str, Any]] = []
        self.format = self.template_path.suffix.lower()
        
    def validate(self) -> Dict[str, Any]:
        """Main validation method"""
        print(f"🔍 Validating template: {self.template_path}")
        print(f"📋 Format: {self.format}")
        
        results = {
            "template": str(self.template_path),
            "format": self.format,
            "valid": True,
            "issues": [],
            "ooxml_structure": {},
            "corruption_risk": "low"
        }
        
        # Basic file validation
        if not self.template_path.exists():
            self.add_issue("critical", "Template file not found")
            results["valid"] = False
            return results
            
        # ZIP structure validation
        zip_valid = self._validate_zip_structure()
        results["zip_valid"] = zip_valid
        
        # OOXML structure validation
        ooxml_structure = self._validate_ooxml_structure()
        results["ooxml_structure"] = ooxml_structure
        
        # Format-specific validation
        if self.format == ".docx" or self.format == ".dotx":
            format_results = self._validate_word_template()
        elif self.format == ".xlsx" or self.format == ".xltx":
            format_results = self._validate_excel_template()
        elif self.format == ".pptx" or self.format == ".potx":
            format_results = self._validate_powerpoint_template()
        else:
            self.add_issue("warning", f"Unknown format: {self.format}")
            format_results = {}
            
        results.update(format_results)
        results["issues"] = self.issues
        results["valid"] = len([i for i in self.issues if i["level"] == "critical"]) == 0
        results["corruption_risk"] = self._assess_corruption_risk()
        
        return results
        
    def add_issue(self, level: str, message: str, details: str = None):
        """Add validation issue"""
        self.issues.append({
            "level": level,
            "message": message,
            "details": details
        })
        
    def _validate_zip_structure(self) -> bool:
        """Validate ZIP archive structure"""
        try:
            with zipfile.ZipFile(self.template_path, "r") as zip_file:
                # Test ZIP integrity
                zip_file.testzip()
                
                # Check required OOXML files
                files = zip_file.namelist()
                
                required_files = ["[Content_Types].xml", "_rels/.rels"]
                missing_files = [f for f in required_files if f not in files]
                
                if missing_files:
                    self.add_issue("critical", f"Missing required files: {missing_files}")
                    return False
                    
                # Check for extension variables
                var_files = [f for f in files if "extensionvariables" in f.lower()]
                if var_files:
                    print(f"📝 Found extension variables: {var_files}")
                    
                return True
                
        except zipfile.BadZipFile:
            self.add_issue("critical", "Template is not a valid ZIP archive")
            return False
        except Exception as e:
            self.add_issue("critical", f"ZIP validation failed: {str(e)}")
            return False
            
    def _validate_ooxml_structure(self) -> Dict[str, Any]:
        """Validate OOXML XML structure"""
        structure = {
            "content_types_valid": False,
            "relationships_valid": False,
            "xml_files_valid": 0,
            "xml_files_total": 0,
            "malformed_xml": []
        }
        
        try:
            with zipfile.ZipFile(self.template_path, "r") as zip_file:
                # Validate [Content_Types].xml
                try:
                    content_types = zip_file.read("[Content_Types].xml")
                    ET.fromstring(content_types)
                    structure["content_types_valid"] = True
                except (ExpatError, XMLSyntaxError) as e:
                    self.add_issue("critical", f"Malformed [Content_Types].xml: {str(e)}")
                except Exception as e:
                    self.add_issue("warning", f"Could not read [Content_Types].xml: {str(e)}")
                    
                # Validate _rels/.rels
                try:
                    rels = zip_file.read("_rels/.rels")
                    ET.fromstring(rels)
                    structure["relationships_valid"] = True
                except (ExpatError, XMLSyntaxError) as e:
                    self.add_issue("critical", f"Malformed _rels/.rels: {str(e)}")
                except Exception as e:
                    self.add_issue("warning", f"Could not read _rels/.rels: {str(e)}")
                    
                # Validate all XML files
                xml_files = [f for f in zip_file.namelist() if f.endswith(".xml")]
                structure["xml_files_total"] = len(xml_files)
                
                for xml_file in xml_files:
                    try:
                        content = zip_file.read(xml_file)
                        ET.fromstring(content)
                        structure["xml_files_valid"] += 1
                    except (ExpatError, XMLSyntaxError) as e:
                        structure["malformed_xml"].append(xml_file)
                        self.add_issue("error", f"Malformed XML in {xml_file}: {str(e)}")
                        
        except Exception as e:
            self.add_issue("error", f"OOXML structure validation failed: {str(e)}")
            
        return structure
        
    def _validate_word_template(self) -> Dict[str, Any]:
        """Validate Word template specific issues"""
        results = {
            "margins_valid": False,
            "styles_count": 0,
            "variables_found": [],
            "corruption_indicators": []
        }
        
        try:
            # For template formats (.dotx), analyze raw OOXML structure
            if self.format == ".dotx":
                results.update(self._analyze_word_template_raw())
            else:
                # Use python-docx for regular documents
                doc = Document(self.template_path)
                
                # Check sections and margins
                if doc.sections:
                    section = doc.sections[0]
                    margins = {
                        "top": section.top_margin.inches if section.top_margin else None,
                        "bottom": section.bottom_margin.inches if section.bottom_margin else None,
                        "left": section.left_margin.inches if section.left_margin else None,
                        "right": section.right_margin.inches if section.right_margin else None
                    }
                    
                    # Check if margins are zero or missing
                    zero_margins = [k for k, v in margins.items() if v == 0 or v is None]
                    if zero_margins:
                        self.add_issue("warning", f"Missing/zero margins detected: {zero_margins}")
                        results["missing_margins"] = zero_margins
                    else:
                        results["margins_valid"] = True
                        
                    results["margins"] = margins
                    
                # Check styles
                try:
                    results["styles_count"] = len(doc.styles)
                    print(f"📊 Found {results["styles_count"]} styles")
                except Exception as e:
                    self.add_issue("warning", f"Could not analyze styles: {str(e)}")
                        
            # Look for extension variables in raw XML
            variables = self._extract_extension_variables()
            results["variables_found"] = variables
            
        except Exception as e:
            self.add_issue("error", f"Word template analysis failed: {str(e)}")
            results["corruption_indicators"].append(f"PyOffice error: {str(e)}")
            
        return results
        
    def _validate_excel_template(self) -> Dict[str, Any]:
        """Validate Excel template specific issues"""
        results = {
            "worksheets_count": 0,
            "corruption_indicators": [],
            "variables_found": []
        }
        
        try:
            # Use openpyxl to analyze
            workbook = load_workbook(self.template_path)
            results["worksheets_count"] = len(workbook.worksheets)
            results["worksheet_names"] = [ws.title for ws in workbook.worksheets]
            
            print(f"📊 Found {results["worksheets_count"]} worksheets: {results["worksheet_names"]}")
            
            # Look for extension variables
            variables = self._extract_extension_variables()
            results["variables_found"] = variables
            
        except Exception as e:
            self.add_issue("error", f"Excel template analysis failed: {str(e)}")
            results["corruption_indicators"].append(f"PyOffice error: {str(e)}")
            
        return results
        
    def _validate_powerpoint_template(self) -> Dict[str, Any]:
        """Validate PowerPoint template specific issues"""
        results = {
            "slides_count": 0,
            "layouts_count": 0,
            "corruption_indicators": [],
            "variables_found": []
        }
        
        try:
            # Use python-pptx to analyze
            presentation = Presentation(self.template_path)
            results["slides_count"] = len(presentation.slides)
            results["layouts_count"] = len(presentation.slide_layouts)
            
            print(f"📊 Found {results["slides_count"]} slides, {results["layouts_count"]} layouts")
            
            # Look for extension variables
            variables = self._extract_extension_variables()
            results["variables_found"] = variables
            
        except Exception as e:
            self.add_issue("error", f"PowerPoint template analysis failed: {str(e)}")
            results["corruption_indicators"].append(f"PyOffice error: {str(e)}")
            
        return results
        
    def _extract_extension_variables(self) -> List[str]:
        """Extract extension variables from template"""
        variables = []
        
        try:
            with zipfile.ZipFile(self.template_path, "r") as zip_file:
                # Look for extension variables in various XML files
                xml_files = [f for f in zip_file.namelist() if f.endswith(".xml")]
                
                for xml_file in xml_files:
                    try:
                        content = zip_file.read(xml_file).decode("utf-8")
                        
                        # Look for StyleStack extension variables
                        if "stylestack" in content.lower() or "extensionvariables" in content.lower():
                            # Simple regex to find variables
                            import re
                            var_pattern = r'\{([^}]+)\}'
                            found_vars = re.findall(var_pattern, content)
                            variables.extend(found_vars)
                            
                    except Exception:
                        continue
                        
        except Exception as e:
            self.add_issue("warning", f"Could not extract variables: {str(e)}")
            
        return list(set(variables))  # Remove duplicates
        
    def _analyze_word_template_raw(self) -> Dict[str, Any]:
        """Analyze Word template using raw OOXML structure"""
        results = {
            "margins_valid": False,
            "styles_count": 0,
            "sections_found": 0,
            "corruption_indicators": []
        }
        
        try:
            with zipfile.ZipFile(self.template_path, "r") as zip_file:
                # Analyze document.xml for structure
                if "word/document.xml" in zip_file.namelist():
                    doc_xml = zip_file.read("word/document.xml")
                    root = ET.fromstring(doc_xml)
                    
                    # Look for sections
                    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
                    sections = root.findall(".//w:sectPr", ns)
                    results["sections_found"] = len(sections)
                    
                    if sections:
                        # Check first section for margins
                        first_section = sections[0]
                        
                        # Look for page margins
                        page_mar = first_section.find(".//w:pgMar", ns)
                        if page_mar is not None:
                            margins = {
                                "top": self._twips_to_inches(page_mar.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}top", "1440")),
                                "bottom": self._twips_to_inches(page_mar.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}bottom", "1440")),
                                "left": self._twips_to_inches(page_mar.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}left", "1440")),
                                "right": self._twips_to_inches(page_mar.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}right", "1440"))
                            }
                            
                            # Check if margins are zero or very small
                            zero_margins = [k for k, v in margins.items() if v == 0 or v < 0.1]
                            if zero_margins:
                                self.add_issue("warning", f"Missing/zero margins detected in raw XML: {zero_margins}")
                                results["missing_margins"] = zero_margins
                            else:
                                results["margins_valid"] = True
                                
                            results["margins"] = margins
                        else:
                            self.add_issue("warning", "No page margin settings found in document")
                            results["corruption_indicators"].append("missing_page_margins")
                    else:
                        self.add_issue("error", "No section properties found in document")
                        results["corruption_indicators"].append("missing_section_properties")
                        
                # Analyze styles.xml
                if "word/styles.xml" in zip_file.namelist():
                    styles_xml = zip_file.read("word/styles.xml")
                    styles_root = ET.fromstring(styles_xml)
                    styles = styles_root.findall(".//w:style", ns)
                    results["styles_count"] = len(styles)
                    print(f"📊 Found {results["styles_count"]} styles in raw XML")
                else:
                    self.add_issue("warning", "No styles.xml found in template")
                    results["corruption_indicators"].append("missing_styles")
                    
        except Exception as e:
            self.add_issue("error", f"Raw Word template analysis failed: {str(e)}")
            results["corruption_indicators"].append(f"Raw analysis error: {str(e)}")
            
        return results
        
    def _twips_to_inches(self, twips_str: str) -> float:
        """Convert twips (1/1440 inch) to inches"""
        try:
            twips = int(twips_str)
            return twips / 1440.0
        except (ValueError, TypeError):
            return 0.0
        
    def _assess_corruption_risk(self) -> str:
        """Assess corruption risk based on issues found"""
        critical_count = len([i for i in self.issues if i["level"] == "critical"])
        error_count = len([i for i in self.issues if i["level"] == "error"])
        
        if critical_count > 0:
            return "high"
        elif error_count > 2:
            return "medium"
        elif error_count > 0:
            return "low"
        else:
            return "minimal"
            
    def print_summary(self, results: Dict[str, Any]):
        """Print validation summary"""
        print("\n" + "="*60)
        print("📋 TEMPLATE VALIDATION SUMMARY")
        print("="*60)
        
        status = "✅ VALID" if results["valid"] else "❌ INVALID"
        print(f"Status: {status}")
        print(f"Corruption Risk: {results["corruption_risk"].upper()}")
        
        if results["issues"]:
            print(f"\n🚨 Issues Found: {len(results["issues"])}")
            for issue in results["issues"]:
                level_emoji = {"critical": "🔴", "error": "🟠", "warning": "🟡"}.get(issue["level"], "📝")
                print(f"  {level_emoji} {issue["level"].upper()}: {issue["message"]}")
                if issue.get("details"):
                    print(f"     Details: {issue["details"]}")
                    
        # Format-specific summary
        if self.format in [".docx", ".dotx"]:
            self._print_word_summary(results)
        elif self.format in [".xlsx", ".xltx"]:
            self._print_excel_summary(results)
        elif self.format in [".pptx", ".potx"]:
            self._print_powerpoint_summary(results)
            
        print("\n" + "="*60)
        
    def _print_word_summary(self, results: Dict[str, Any]):
        """Print Word-specific summary"""
        print(f"\n📄 WORD TEMPLATE ANALYSIS:")
        if "margins" in results:
            print(f"   Margins: {results["margins"]}")
        if "missing_margins" in results:
            print(f"   ⚠️  Missing margins: {results["missing_margins"]}")
        print(f"   Styles: {results.get("styles_count", "unknown")}")
        if results.get("variables_found"):
            print(f"   Variables: {len(results["variables_found"])} found")
            
    def _print_excel_summary(self, results: Dict[str, Any]):
        """Print Excel-specific summary"""
        print(f"\n📊 EXCEL TEMPLATE ANALYSIS:")
        print(f"   Worksheets: {results.get("worksheets_count", 0)}")
        if results.get("worksheet_names"):
            print(f"   Names: {results["worksheet_names"]}")
        if results.get("variables_found"):
            print(f"   Variables: {len(results["variables_found"])} found")
            
    def _print_powerpoint_summary(self, results: Dict[str, Any]):
        """Print PowerPoint-specific summary"""
        print(f"\n🎯 POWERPOINT TEMPLATE ANALYSIS:")
        print(f"   Slides: {results.get("slides_count", 0)}")
        print(f"   Layouts: {results.get("layouts_count", 0)}")
        if results.get("variables_found"):
            print(f"   Variables: {len(results["variables_found"])} found")


def main():
    """Main CLI entry point"""
    parser = argparse.ArgumentParser(description="Validate StyleStack templates")
    parser.add_argument("--template", "-t", required=True,
                       help="Path to template file (.docx/.dotx/.xlsx/.xltx/.pptx/.potx)")
    parser.add_argument("--verbose", "-v", action="store_true",
                       help="Verbose output")
    
    args = parser.parse_args()
    
    # Reminder about venv
    print("🔍 StyleStack Template Validator")
    print("⚠️  IMPORTANT: Ensure venv is activated (source venv/bin/activate)")
    print()
    
    validator = TemplateValidator(args.template)
    results = validator.validate()
    validator.print_summary(results)
    
    # Exit with error code if validation failed
    sys.exit(0 if results["valid"] else 1)


if __name__ == "__main__":
    main()